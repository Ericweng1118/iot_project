"""
data_layer/import_templates.py
==============================
「批次匯入匯出」的 Excel 匯入範本。

匯出現有資料再改適合「修改」，要「從零建立」時那份檔案反而礙事（要先刪光資料列、
不知道哪些欄位必填、值該填什麼）。範本提供：
    資料    第一個工作表，標題列跟匯入格式完全一致；必填欄位標橘色，滑鼠移到標題上有說明；
            選項固定的欄位（上傳條件、功能碼…）與參照欄位（設備、感測器）有下拉選單
    說明    每個欄位的必填 / 說明
    選單    下拉選單的來源（隱藏），由目前資料庫內容產生

感測器下拉選單顯示「編號｜設備｜暱稱」，匯入時 bulk_io.p_sensor_ref 只取「｜」前的編號。
OPC UA 綁定的點位只能由瀏覽產生，範本會預先列出所有「尚未綁定」的點位，只要填 sensor_code。
"""

import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from data_layer.bulk_io import (
    ALARM_TYPES,
    DEADBAND_TYPES,
    MODBUS_COLUMNS,
    S7_COLUMNS,
    SENSOR_REF_SEP,
    TRANSPORTS,
    UPLOAD_CONDITIONS,
    _has_table,
    _query,
)
from protocols import s7_codec
from protocols.modbus_codec import DATA_TYPES as MODBUS_TYPES

# 下拉選單至少套用到第幾列（預先列出的資料比這多時會自動延伸）
DATA_ROWS = 1000

COLUMNS = {
    "devices": ["device_code", "device_name", "device_type", "manufacturer", "site_name", "line_name", "status"],
    "sensors": ["sensor_code", "device_code", "sensor_type", "nickname", "unit", "min_threshold", "max_threshold",
                "upload_condition", "upload_threshold", "opcua_sampling_interval_ms", "opcua_deadband_type",
                "opcua_deadband_value", "state_dictionary"],
    "modbus": MODBUS_COLUMNS,
    "s7": S7_COLUMNS,
    "opcua_bindings": ["server_name", "node_id", "display_name", "data_type", "sensor_code"],
    "alarm_rules": ["rule_id", "sensor_code", "alarm_type", "setpoint", "deadband", "on_delay_sec", "priority",
                    "message", "enabled"],
}

# 與 bulk_io.plan_* 的必填欄位一致（標橘色用）
REQUIRED = {
    "devices": {"device_code", "site_name", "line_name"},
    "sensors": {"device_code", "sensor_type"},
    "modbus": {"name", "plc_ip", "slave_id", "function_code", "start_address", "data_type"},
    "s7": {"name", "plc_ip", "data_type"},
    "opcua_bindings": {"server_name", "node_id"},
    "alarm_rules": {"sensor_code", "alarm_type", "setpoint"},
}

YES_NO = ["是", "否"]
SENSOR_TYPES = ["temperature", "pressure", "vibration", "current", "voltage", "power", "energy", "flow",
                "level", "humidity", "speed", "torque", "position", "ph", "conductivity", "weight", "count",
                "status", "gas", "liquid", "solid"]

_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
_REQUIRED_FILL = PatternFill("solid", fgColor="F8CBAD")


def _distinct(cur, sql) -> list:
    df = _query(cur, sql)
    return [] if df.empty else [str(v) for v in df.iloc[:, 0] if pd.notna(v) and str(v).strip()]


def _merge(*lists) -> list:
    out = []
    for lst in lists:
        for v in lst:
            if v not in out:
                out.append(v)
    return out


def sensor_choices(cur, free_only: bool) -> list:
    """「編號｜設備｜暱稱」清單；free_only=True 時排除已被點位或計算點使用的感測器。"""
    used = ""
    if free_only:
        parts = [f"SELECT sensor_id FROM {t} WHERE sensor_id IS NOT NULL"
                 for t in ("modbus_scada", "tia_scada", "opcua_tags")]
        if _has_table(cur, "calculated_points"):
            parts.append("SELECT sensor_id FROM calculated_points")
        used = f"WHERE s.sensor_id NOT IN ({' UNION '.join(parts)})"
    df = _query(cur, f"""
        SELECT s.sensor_code, d.device_code, s.nickname
        FROM sensors s LEFT JOIN devices d ON d.device_id = s.device_id
        {used}
        ORDER BY d.device_code NULLS LAST, length(s.sensor_code), s.sensor_code;""")
    return [SENSOR_REF_SEP.join(str(v) for v in (r["sensor_code"], r["device_code"], r["nickname"])
                                if pd.notna(v) and str(v).strip())
            for _, r in df.iterrows()]


def _choices(cur, entity) -> dict:
    """{欄位: (選項, 是否只能選清單內的值)}；非嚴格的只是提示，仍可自行輸入。"""
    if entity == "devices":
        return {
            "site_name": (_distinct(cur, "SELECT site_name FROM sites ORDER BY 1;"), False),
            "line_name": (_distinct(cur, "SELECT DISTINCT line_name FROM production_lines ORDER BY 1;"), False),
            "status": (_merge(["active", "inactive"],
                              _distinct(cur, "SELECT DISTINCT status FROM devices ORDER BY 1;")), False),
        }
    if entity == "sensors":
        return {
            "device_code": (_distinct(cur, "SELECT device_code FROM devices ORDER BY 1;"), True),
            "sensor_type": (_merge(SENSOR_TYPES,
                                   _distinct(cur, "SELECT DISTINCT sensor_type FROM sensors ORDER BY 1;")), False),
            "unit": (_distinct(cur, "SELECT DISTINCT unit FROM sensors ORDER BY 1;"), False),
            "upload_condition": (list(UPLOAD_CONDITIONS), True),
            "opcua_deadband_type": (list(DEADBAND_TYPES), True),
        }
    if entity == "modbus":
        return {
            "transport": (list(TRANSPORTS), True),
            "function_code": ([1, 2, 3, 4], True),
            "data_type": (list(MODBUS_TYPES), False),
            "byte_order": (["BIG", "LITTLE"], True),
            "word_order": (["BIG", "LITTLE"], True),
            "sensor_code": (sensor_choices(cur, free_only=True), False),
            "enabled": (YES_NO, False),
        }
    if entity == "s7":
        return {
            "area": (["DB", "M", "I", "Q"], True),
            "data_type": (list(s7_codec.TYPES), False),
            "sensor_code": (sensor_choices(cur, free_only=True), False),
            "enabled": (YES_NO, False),
        }
    if entity == "opcua_bindings":
        return {"sensor_code": (sensor_choices(cur, free_only=True), False)}
    if entity == "alarm_rules":
        return {
            "sensor_code": (sensor_choices(cur, free_only=False), False),
            "alarm_type": (list(ALARM_TYPES), True),
            "priority": ([1, 2, 3, 4], True),
            "enabled": (YES_NO, False),
        }
    return {}


def _prefill(cur, entity) -> pd.DataFrame | None:
    if entity == "opcua_bindings":
        return _query(cur, """
            SELECT server_name, node_id, display_name, data_type, NULL AS sensor_code
            FROM opcua_tags WHERE sensor_id IS NULL
            ORDER BY server_name, node_id;""")
    return None


def _doc_for(col: str, docs) -> str:
    """從頁面上的欄位說明（可能是「a / b」合併寫法）找出這個欄位的說明文字。"""
    for names, required, text in docs:
        parts = [p.strip() for p in str(names).split("/")]
        if col in parts or any(p.startswith("_") and col.endswith(p) for p in parts):
            return f"{'【必填】' if required.strip() == '✅' else ''}{text}"
    return ""


def build_template(cur, entity: str, docs=()) -> bytes:
    """
    :param docs: 頁面上的欄位說明 [(欄位, 必填標記, 說明)]，寫進「說明」工作表與標題註解
    :return: xlsx bytes
    """
    columns = COLUMNS[entity]
    required = REQUIRED.get(entity, set())
    choices = _choices(cur, entity)
    prefill = _prefill(cur, entity)

    wb = Workbook()
    ws = wb.active
    ws.title = "資料"            # 匯入時讀第一個工作表，一定要放第一個
    ws.append(columns)
    ws.freeze_panes = "A2"
    for idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True)
        cell.fill = _REQUIRED_FILL if col in required else _HEADER_FILL
        doc = _doc_for(col, docs)
        if doc:
            cell.comment = Comment(doc, "SCADA")
        ws.column_dimensions[get_column_letter(idx)].width = max(14, min(42, len(col) + 6))
    if "sensor_code" in choices:
        ws.column_dimensions[get_column_letter(columns.index("sensor_code") + 1)].width = 42

    rows = DATA_ROWS
    if prefill is not None and not prefill.empty:
        for rec in prefill[columns].itertuples(index=False):
            ws.append([None if pd.isna(v) else v for v in rec])
        rows = max(rows, len(prefill) + 200)

    # 下拉選單來源放在隱藏的「選單」工作表，一個欄位一欄
    lists = wb.create_sheet("選單")
    list_col = 0
    for col, (values, strict) in choices.items():
        if col not in columns or not values:
            continue
        list_col += 1
        letter = get_column_letter(list_col)
        lists.cell(row=1, column=list_col, value=col)
        for r, v in enumerate(values, start=2):
            lists.cell(row=r, column=list_col, value=v)
        dv = DataValidation(
            type="list", formula1=f"'選單'!${letter}$2:${letter}${len(values) + 1}",
            allow_blank=True, showErrorMessage=strict,
            errorTitle="不在選項內", error=f"{col} 只能從下拉選單選擇",
        )
        target = get_column_letter(columns.index(col) + 1)
        dv.add(f"{target}2:{target}{rows + 1}")
        ws.add_data_validation(dv)
    lists.sheet_state = "hidden"

    notes = wb.create_sheet("說明", 1)
    notes.append(["使用方式"])
    notes["A1"].font = Font(bold=True)
    for line in (
        "1. 在「資料」工作表從第 2 列開始填，一列一筆；不要修改或刪除標題列。",
        "2. 橘色標題是必填欄位；滑鼠移到標題上可以看說明。有下拉選單的欄位建議直接用選的。",
        f"3. 感測器下拉選單是「編號{SENSOR_REF_SEP}設備{SENSOR_REF_SEP}暱稱」，匯入時只取第一段的編號，"
        "直接手打編號也可以。",
        "4. 存檔後到網頁「批次匯入匯出」上傳，會先顯示預覽（新增 / 更新 / 錯誤），確認後才寫入。",
    ):
        notes.append([line])
    if entity == "sensors":
        notes.append(["5. sensor_code 留空 = 新增感測器並自動編號（建議）；填既有編號 = 更新那個感測器。"])
    if entity == "opcua_bindings":
        notes.append(["5. 已預先列出所有尚未綁定的點位，只要填 sensor_code；不需要的列可以刪掉或留空（留空 = 不綁定）。"])
    notes.append([])
    notes.append(["欄位", "必填", "說明"])
    for c in notes[notes.max_row]:
        c.font = Font(bold=True)
        c.fill = _HEADER_FILL
    for names, req, text in docs:
        notes.append([names, req, text])
    notes.column_dimensions["A"].width = 38
    notes.column_dimensions["B"].width = 10
    notes.column_dimensions["C"].width = 90
    for row in notes.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=False, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
